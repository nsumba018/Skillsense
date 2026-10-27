"""
Parse Excel bulletin files into structured course data and extract marketable skills.

The parser handles AUCA-style bulletins with sheets per department, each containing:
  Semester | CourseName | Credits
"""
import re
from pathlib import Path
from django.core.files.uploadedfile import UploadedFile
import openpyxl

# ── Course-name → marketable-skill mapping ──────────────────────────────
# Each key is a regex (case-insensitive) matched against course names.
# Values are the list of marketable skills that course teaches.
COURSE_SKILL_MAP: list[tuple[str, list[str]]] = [
    # Programming languages
    (r'programming\s*with\s*c\b', ['C', 'Programming']),
    (r'intro.*computer programming', ['Programming', 'Problem Solving']),
    (r'java\s*programming', ['Java', 'Java/SpringBoot', 'OOP']),
    (r'object.oriented\s*programming', ['OOP', 'Java', 'C#']),
    (r'dot\s*net|\.net', ['.NET', 'C#']),
    (r'pl/?sql', ['PL/SQL', 'SQL', 'Database Design']),
    (r'python|big\s*data', ['Python', 'Data Analysis', 'Big Data']),
    (r'best\s*programming.*design\s*patterns', ['Design Patterns', 'Clean Code', 'Architecture']),

    # Web & Mobile
    (r'web\s*design', ['HTML/CSS', 'Web Design', 'UI/UX']),
    (r'web\s*technology', ['JavaScript', 'React', 'HTML/CSS', 'REST APIs', 'Web Development']),
    (r'mobile\s*programming|mobile\s*communication', ['Mobile Development', 'Android', 'iOS']),
    (r'e-commerce', ['E-Commerce', 'Web Development']),

    # Software engineering
    (r'^software\s*engineering$', ['Software Engineering', 'SDLC', 'Agile/Scrum']),
    (r'requirements\s*engineering', ['Requirements Engineering', 'Business Analysis']),
    (r'software\s*modeling|software\s*design', ['UML', 'Software Modeling', 'System Design']),
    (r'software\s*testing', ['Software Testing', 'QA', 'Test Automation']),
    (r'software\s*quality', ['QA', 'Software Quality Assurance', 'Testing']),
    (r'software\s*project\s*management', ['Project Management', 'Agile/Scrum', 'Jira']),
    (r'software\s*security', ['Software Security', 'Cybersecurity', 'OWASP']),
    (r'system\s*analysis.*design', ['System Analysis', 'System Design', 'UML']),

    # Data & databases
    (r'database\s*management|dbms', ['DBMS', 'SQL', 'Database Design']),
    (r'data\s*base\s*development|pl/?sql', ['PL/SQL', 'SQL', 'Database Design']),
    (r'data\s*structure', ['Data Structures', 'Algorithms']),
    (r'introduction\s*to\s*big\s*data', ['Python', 'Big Data', 'Data Analysis', 'Machine Learning']),

    # Networking & infrastructure
    (r'computer\s*networks$', ['Computer Networks', 'TCP/IP', 'Networking']),
    (r'advanced\s*computer\s*networks', ['Advanced Networking', 'TCP/IP', 'Network Architecture']),
    (r'routing\s*and\s*switching', ['Cisco', 'Routing', 'Switching', 'CCNA']),
    (r'network\s*security', ['Network Security', 'Firewall', 'Cybersecurity', 'Penetration Testing']),
    (r'network\s*admin', ['Network Administration', 'Windows Server', 'Active Directory']),
    (r'network\s*programming', ['TCP/IP', 'Socket Programming', 'Network Programming']),
    (r'wireless\s*networks', ['Wireless Networking', 'Wi-Fi', 'RF']),
    (r'system\s*admin', ['System Administration', 'Linux', 'Windows Server']),

    # Operating systems & linux
    (r'operating\s*systems', ['Operating Systems', 'Linux', 'Process Management']),
    (r'unix|linux', ['Linux', 'Shell Scripting', 'System Administration']),

    # Math & statistics
    (r'statistics\s*and\s*probability|descriptive\s*statistics', ['Statistics', 'Probability', 'Data Analysis']),
    (r'applied\s*math|calculus|multivariable', ['Mathematics', 'Calculus']),
    (r'theory\s*of\s*computation', ['Theory of Computation', 'Discrete Mathematics']),

    # Management & business
    (r'management\s*information|^mis$', ['MIS', 'Information Systems', 'Business Analysis']),
    (r'principles\s*of\s*management', ['Management', 'Leadership', 'Organizational Behavior']),
    (r'strategic\s*management', ['Strategic Management', 'Business Strategy']),
    (r'human\s*resource|hr\s*management', ['HR Management', 'People Management']),
    (r'marketing', ['Marketing', 'Market Research']),
    (r'entrepreneurship', ['Entrepreneurship', 'Business Development']),
    (r'operations\s*management', ['Operations', 'Supply Chain']),

    # Finance & accounting
    (r'principles\s*of\s*accounting', ['Principles of Accounting', 'Financial Accounting']),
    (r'intermediate\s*accounting', ['Financial Accounting', 'Auditing', 'IFRS']),
    (r'advanced\s*accounting', ['Advanced Accounting', 'IFRS', 'Corporate Finance']),
    (r'cost\s*accounting', ['Cost Accounting', 'Budgeting', 'Financial Analysis']),
    (r'corporate\s*finance', ['Corporate Finance', 'Investment', 'Financial Analysis']),
    (r'banking', ['Banking', 'Financial Services']),
    (r'tax', ['Tax', 'Compliance']),
    (r'economics', ['Economics', 'Microeconomics', 'Macroeconomics']),
    (r'business\s*law', ['Business Law', 'Compliance']),

    # Other
    (r'digital\s*computer\s*fund', ['Digital Logic', 'Computer Architecture']),
    (r'computer\s*maintenance', ['Hardware Troubleshooting', 'IT Support']),
    (r'multimedia', ['Multimedia', 'Content Creation']),
    (r'emerging\s*tech|emmerging\s*tech', ['Emerging Technologies', 'IoT', 'AI/ML']),
    (r'information\s*security', ['Information Security', 'Cybersecurity', 'Risk Management']),

    # General / communication
    (r'english|academic.*writing', ['Communication', 'English', 'Technical Writing']),
    (r'study.*research', ['Research Methods', 'Academic Writing']),
]

# Department-code detection from sheet or header names
DEPT_PATTERNS: list[tuple[str, str, str, str]] = [
    # (regex on sheet/header, department_code, department_name, faculty)
    (r'software', 'se', 'Software Engineering', 'it'),
    (r'network', 'net', 'Networks & Communication', 'it'),
    (r'information.m', 'im', 'Information Management', 'it'),
    (r'big\s*data|data\s*analytics', 'bda', 'Big Data Analytics', 'it'),
    (r'accounting', 'acc', 'Accounting', 'business'),
    (r'finance', 'fin', 'Finance', 'business'),
    (r'management(?!.*info)', 'ba', 'Management', 'business'),
    (r'marketing', 'mkt', 'Marketing', 'business'),
    (r'mba', 'mba', 'MBA', 'business'),
]


def _detect_department(text: str) -> tuple[str, str, str]:
    """Return (code, name, faculty) from a sheet name or header text."""
    low = text.lower()
    for pattern, code, name, faculty in DEPT_PATTERNS:
        if re.search(pattern, low):
            return code, name, faculty
    return 'unknown', text.strip(), 'it'


def _extract_skills(course_name: str) -> list[str]:
    """Map a course name to marketable skills."""
    clean = course_name.replace('\xa0', ' ').strip()
    skills: list[str] = []
    for pattern, mapped in COURSE_SKILL_MAP:
        if re.search(pattern, clean, re.IGNORECASE):
            skills.extend(mapped)
    return list(dict.fromkeys(skills))  # dedupe, preserve order


def parse_excel(file_path: str | Path) -> list[dict]:
    """
    Parse an Excel bulletin file.

    Returns a list of dicts, one per sheet/department:
      {
        'department_code': str,
        'department_name': str,
        'faculty': str,
        'total_credits': int,
        'total_semesters': int,
        'courses': [{'semester': int, 'name': str, 'credits': int, 'skills': [...]}]
      }
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    results = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        dept_code, dept_name, faculty = _detect_department(sheet_name)

        # Also try to detect from the header rows
        for row in ws.iter_rows(min_row=1, max_row=3, values_only=True):
            for cell in row:
                if cell and isinstance(cell, str) and 'department' in cell.lower():
                    c, n, f = _detect_department(cell)
                    if c != 'unknown':
                        dept_code, dept_name, faculty = c, n, f

        courses = []
        semesters_seen = set()

        for row in ws.iter_rows(min_row=1, values_only=True):
            vals = [v for v in row]
            if len(vals) < 3:
                continue

            sem_raw, name_raw, credits_raw = vals[0], vals[1], vals[2]

            # Skip header/total/empty rows
            if sem_raw is None or name_raw is None:
                continue
            if isinstance(sem_raw, str) and sem_raw.strip().upper() in ('SEMESTER', 'TOTAL', 'GRAND TOTAL', ''):
                continue
            if isinstance(credits_raw, str) and credits_raw.startswith('='):
                continue

            try:
                semester = int(sem_raw)
            except (ValueError, TypeError):
                continue

            name = str(name_raw).replace('\xa0', ' ').strip()
            if not name:
                continue

            try:
                credits = int(credits_raw) if credits_raw else 0
            except (ValueError, TypeError):
                credits = 0

            semesters_seen.add(semester)
            skills = _extract_skills(name)
            courses.append({
                'semester': semester,
                'name': name,
                'credits': credits,
                'skills': skills,
            })

        total_credits = sum(c['credits'] for c in courses)

        results.append({
            'department_code': dept_code,
            'department_name': dept_name,
            'faculty': faculty,
            'total_credits': total_credits,
            'total_semesters': len(semesters_seen),
            'courses': courses,
        })

    return results


def parse_uploaded_file(uploaded: UploadedFile) -> list[dict]:
    """Parse an in-memory uploaded file by saving to a temp path first."""
    import tempfile, os
    suffix = Path(uploaded.name).suffix
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        for chunk in uploaded.chunks():
            tmp.write(chunk)
        tmp_path = tmp.name
    try:
        return parse_excel(tmp_path)
    finally:
        os.unlink(tmp_path)
